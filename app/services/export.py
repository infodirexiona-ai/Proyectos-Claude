"""Exportación de documentos y reportes a Excel y CSV."""

from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..sii import endpoints as ep
from . import reports
from .periodos import periodo_legible

_FORMATO_MONEDA = "#,##0"
_RELLENO_CABECERA = PatternFill("solid", fgColor="1F3864")
_FUENTE_CABECERA = Font(color="FFFFFF", bold=True)

COLUMNAS_DOCUMENTOS = [
    ("Periodo", "periodo", None),
    ("Operación", "operacion", None),
    ("Tipo doc.", "tipo_doc", None),
    ("Documento", "tipo_doc_nombre", None),
    ("Folio", "folio", None),
    ("RUT contraparte", "rut_contraparte", None),
    ("Razón social", "razon_social", None),
    ("Fecha emisión", "fecha_emision", "dd-mm-yyyy"),
    ("Fecha recepción", "fecha_recepcion", "dd-mm-yyyy hh:mm"),
    ("Exento", "monto_exento", _FORMATO_MONEDA),
    ("Neto", "monto_neto", _FORMATO_MONEDA),
    ("IVA", "monto_iva", _FORMATO_MONEDA),
    ("Total", "monto_total", _FORMATO_MONEDA),
    ("IVA no recuperable", "iva_no_recuperable", _FORMATO_MONEDA),
    ("IVA uso común", "iva_uso_comun", _FORMATO_MONEDA),
    ("Otros impuestos", "otros_impuestos", _FORMATO_MONEDA),
    ("Estado", "estado_contab", None),
    ("Glosa (resumen)", "glosa_resumida", None),
    ("Glosa (detalle)", "glosa_extendida", None),
]


def _valor(documento, campo):
    valor = getattr(documento, campo, "")
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None)
    return valor


def _escribir_hoja(libro: Workbook, titulo: str, cabeceras: list[str], filas: list[list], formatos=None):
    hoja = libro.create_sheet(titulo[:31])
    hoja.append(cabeceras)
    for celda in hoja[1]:
        celda.fill = _RELLENO_CABECERA
        celda.font = _FUENTE_CABECERA
        celda.alignment = Alignment(horizontal="center", vertical="center")
    for fila in filas:
        hoja.append(fila)
    for indice, cabecera in enumerate(cabeceras, start=1):
        letra = get_column_letter(indice)
        ancho = max(len(str(cabecera)) + 2, 12)
        for fila in filas[:200]:
            valor = fila[indice - 1] if indice - 1 < len(fila) else ""
            ancho = max(ancho, min(len(str(valor)) + 2, 45))
        hoja.column_dimensions[letra].width = ancho
        if formatos and formatos.get(indice - 1):
            for celda in hoja[letra][1:]:
                celda.number_format = formatos[indice - 1]
    if filas:
        hoja.auto_filter.ref = hoja.dimensions
    hoja.freeze_panes = "A2"
    return hoja


def excel_documentos(db: Session, filtro: reports.Filtro, *, limite: int = 50_000) -> io.BytesIO:
    """Libro Excel con el detalle y las hojas de reportes."""
    documentos, _ = reports.listar_documentos(db, filtro, limite=limite)

    libro = Workbook()
    libro.remove(libro.active)

    _escribir_hoja(
        libro,
        "Documentos",
        [c[0] for c in COLUMNAS_DOCUMENTOS],
        [[_valor(d, c[1]) for c in COLUMNAS_DOCUMENTOS] for d in documentos],
        {i: c[2] for i, c in enumerate(COLUMNAS_DOCUMENTOS) if c[2]},
    )

    _escribir_hoja(
        libro,
        "Por periodo",
        ["Periodo", "Operación", "Documentos", "Exento", "Neto", "IVA", "Total"],
        [
            [
                f["periodo_legible"],
                f["operacion"],
                f["documentos"],
                f["exento"],
                f["neto"],
                f["iva"],
                f["total"],
            ]
            for f in reports.por_periodo(db, filtro)
        ],
        {3: _FORMATO_MONEDA, 4: _FORMATO_MONEDA, 5: _FORMATO_MONEDA, 6: _FORMATO_MONEDA},
    )

    _escribir_hoja(
        libro,
        "Por tipo de documento",
        ["Código", "Documento", "Descripción", "Operación", "Documentos", "Neto", "IVA", "Total"],
        [
            [
                f["tipo_doc"],
                f["tipo_doc_nombre"],
                ep.descripcion_tipo_dte(f["tipo_doc"]),
                f["operacion"],
                f["documentos"],
                f["neto"],
                f["iva"],
                f["total"],
            ]
            for f in reports.por_tipo_documento(db, filtro)
        ],
        {5: _FORMATO_MONEDA, 6: _FORMATO_MONEDA, 7: _FORMATO_MONEDA},
    )

    _escribir_hoja(
        libro,
        "Por contraparte",
        ["RUT", "Razón social", "Operación", "Documentos", "Neto", "IVA", "Total"],
        [
            [f["rut"], f["razon_social"], f["operacion"], f["documentos"], f["neto"], f["iva"], f["total"]]
            for f in reports.por_contraparte(db, filtro, limite=500)
        ],
        {4: _FORMATO_MONEDA, 5: _FORMATO_MONEDA, 6: _FORMATO_MONEDA},
    )

    _escribir_hoja(
        libro,
        "IVA por periodo",
        ["Periodo", "Ventas netas", "Compras netas", "IVA débito", "IVA crédito", "Saldo"],
        [
            [
                f["periodo_legible"],
                f["ventas_netas"],
                f["compras_netas"],
                f["iva_debito"],
                f["iva_credito"],
                f["saldo"],
            ]
            for f in reports.resumen_iva(db, filtro)
        ],
        {1: _FORMATO_MONEDA, 2: _FORMATO_MONEDA, 3: _FORMATO_MONEDA, 4: _FORMATO_MONEDA, 5: _FORMATO_MONEDA},
    )

    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


def csv_documentos(db: Session, filtro: reports.Filtro, *, limite: int = 100_000) -> io.StringIO:
    """CSV plano del detalle (delimitado por ``;``, como los del SII)."""
    documentos, _ = reports.listar_documentos(db, filtro, limite=limite)
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=";", lineterminator="\n")
    escritor.writerow([c[0] for c in COLUMNAS_DOCUMENTOS])
    for documento in documentos:
        escritor.writerow(["" if (v := _valor(documento, c[1])) is None else v for c in COLUMNAS_DOCUMENTOS])
    buffer.seek(0)
    return buffer


def nombre_archivo(filtro: reports.Filtro, extension: str) -> str:
    partes = ["rcv", filtro.rut_titular.replace(".", "")]
    if filtro.operacion:
        partes.append(filtro.operacion.lower())
    if filtro.periodo_desde:
        partes.append(filtro.periodo_desde)
    if filtro.periodo_hasta and filtro.periodo_hasta != filtro.periodo_desde:
        partes.append(filtro.periodo_hasta)
    return "_".join(partes) + f".{extension}"


__all__ = [
    "excel_documentos",
    "csv_documentos",
    "nombre_archivo",
    "periodo_legible",
    "ep",
]
