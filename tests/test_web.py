"""Pruebas de extremo a extremo de la interfaz, en modo demo."""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

RUT = "76192083-9"


@pytest.fixture
def cliente_con_datos(cliente):
    respuesta = cliente.post(
        "/sincronizar",
        data={"rut": RUT, "desde": "202403", "hasta": "202404", "compras": "on", "ventas": "on"},
        follow_redirects=False,
    )
    assert respuesta.status_code == 303
    return cliente


def test_salud(cliente):
    datos = cliente.get("/api/salud").json()
    assert datos["estado"] == "ok"
    assert datos["modo"] == "demo"


def test_panel_sin_datos(cliente):
    respuesta = cliente.get("/")
    assert respuesta.status_code == 200
    assert "Descargar del SII" in respuesta.text


def test_sincronizacion_demo_carga_documentos(cliente_con_datos):
    trabajo = cliente_con_datos.get("/api/sincronizaciones/1").json()
    assert trabajo["estado"] == "ok"
    assert trabajo["nuevos"] > 0

    datos = cliente_con_datos.get("/api/documentos", params={"titular": RUT}).json()
    assert datos["total"] > 0
    primero = datos["documentos"][0]
    assert primero["periodo"] in {"202403", "202404"}
    assert primero["operacion"] in {"COMPRA", "VENTA"}
    assert primero["tipo_doc_nombre"]


def test_paginas_principales(cliente_con_datos):
    for ruta in ("/", "/documentos", "/reportes", "/sincronizaciones", "/credenciales"):
        respuesta = cliente_con_datos.get(ruta)
        assert respuesta.status_code == 200, ruta
        assert "<table" in respuesta.text or "formulario-linea" in respuesta.text


def test_filtros_de_documentos(cliente_con_datos):
    respuesta = cliente_con_datos.get(
        "/documentos", params={"titular": RUT, "operacion": "VENTA", "desde": "202403", "hasta": "202403"}
    )
    assert respuesta.status_code == 200
    assert "COMPRA" not in respuesta.text.split("<tbody>")[1]


def test_reporte_periodo_api(cliente_con_datos):
    datos = cliente_con_datos.get("/api/reportes/periodo", params={"titular": RUT}).json()
    assert len(datos["por_periodo"]) == 4  # 2 periodos x 2 operaciones
    assert len(datos["iva"]) == 2
    assert all("saldo" in fila for fila in datos["iva"])


def test_exportar_excel(cliente_con_datos):
    respuesta = cliente_con_datos.get("/exportar.xlsx", params={"titular": RUT})
    assert respuesta.status_code == 200
    assert "spreadsheetml" in respuesta.headers["content-type"]
    libro = load_workbook(io.BytesIO(respuesta.content))
    assert libro.sheetnames == [
        "Documentos",
        "Por periodo",
        "Por tipo de documento",
        "Por contraparte",
        "IVA por periodo",
    ]
    hoja = libro["Documentos"]
    assert hoja.max_row > 1
    assert hoja["A1"].value == "Periodo"


def test_exportar_csv(cliente_con_datos):
    respuesta = cliente_con_datos.get("/exportar.csv", params={"titular": RUT, "operacion": "COMPRA"})
    assert respuesta.status_code == 200
    lineas = respuesta.text.strip().splitlines()
    assert lineas[0].startswith("﻿Periodo;Operación;")
    assert len(lineas) > 1
    assert all(";COMPRA;" in linea for linea in lineas[1:])


def test_credenciales_ciclo_completo(cliente):
    cliente.post("/credenciales", data={"rut": RUT, "clave": "secreta123", "alias": "Mi empresa"})
    pagina = cliente.get("/credenciales")
    assert RUT in pagina.text
    assert "secreta123" not in pagina.text

    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import Credencial

    with SessionLocal() as db:
        credencial = db.scalars(select(Credencial)).one()
        assert credencial.clave_cifrada != "secreta123"
        cred_id = credencial.id

    cliente.post(f"/credenciales/{cred_id}/eliminar")
    with SessionLocal() as db:
        assert db.scalars(select(Credencial)).all() == []


def test_reenviar_credencial_con_el_mismo_rut_actualiza_en_vez_de_duplicar(cliente):
    """ "Editar" precarga el formulario con el mismo RUT; guardarlo de nuevo
    debe reemplazar la clave, no crear una segunda credencial."""
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import Credencial

    cliente.post("/credenciales", data={"rut": RUT, "clave": "clave-vieja-mal-escrita"})
    cliente.post("/credenciales", data={"rut": RUT, "clave": "clave-correcta"})

    with SessionLocal() as db:
        credenciales = db.scalars(select(Credencial)).all()
        assert len(credenciales) == 1

        from app.config import get_settings
        from app.crypto import descifrar

        clave_guardada = descifrar(credenciales[0].clave_cifrada, get_settings().clave_cifrado)
        assert clave_guardada == "clave-correcta"


def test_boton_editar_trae_los_datos_de_la_fila(cliente):
    cliente.post("/credenciales", data={"rut": RUT, "clave": "algo", "alias": "Mi empresa"})
    pagina = cliente.get("/credenciales")
    assert 'data-rut-titular="76192083-9"' in pagina.text
    assert 'data-alias="Mi empresa"' in pagina.text


def test_rut_invalido_rechazado(cliente):
    respuesta = cliente.post("/sincronizar", data={"rut": "12345678-9", "desde": "202403", "hasta": "202403"})
    assert respuesta.status_code == 400


def test_rango_demasiado_largo(cliente):
    respuesta = cliente.post("/sincronizar", data={"rut": RUT, "desde": "201901", "hasta": "202412"})
    assert respuesta.status_code == 400
    assert "36 periodos" in respuesta.json()["detail"]


def test_descarga_fuera_del_directorio_rechazada(cliente):
    respuesta = cliente.get("/descargas/../../etc/passwd")
    assert respuesta.status_code in (400, 404)


def test_formato_clp_pone_el_signo_antes_del_simbolo():
    from app.web.formato import clp

    assert clp(1234567) == "$1.234.567"
    assert clp(-958360) == "-$958.360"
    assert clp(0) == "$0"
    assert clp(None) == "—"


def test_diagnostico_pagina(cliente):
    respuesta = cliente.get("/diagnostico?red=false")
    assert respuesta.status_code == 200
    assert "Diagnóstico" in respuesta.text
    assert "Todo listo" in respuesta.text  # modo demo: no exige navegador ni red


def test_diagnostico_api(cliente):
    datos = cliente.get("/api/diagnostico?red=false").json()
    assert datos["listo"] is True
    nombres = {c["nombre"] for c in datos["chequeos"]}
    assert "Modo de operación" in nombres
    assert "Base de datos" in nombres


def test_detalle_ventas_demo_completa_la_glosa(cliente_con_datos):
    respuesta = cliente_con_datos.post(
        "/sincronizar-detalle-ventas",
        data={"rut": RUT, "desde": "202403", "hasta": "202404"},
        follow_redirects=False,
    )
    assert respuesta.status_code == 303

    trabajo = cliente_con_datos.get("/api/sincronizaciones/2").json()
    assert trabajo["estado"] == "ok"
    assert trabajo["actualizados"] > 0

    pagina = cliente_con_datos.get("/documentos", params={"titular": RUT, "operacion": "VENTA"})
    assert "Ver (" in pagina.text


def test_detalle_ventas_rut_invalido(cliente):
    respuesta = cliente.post(
        "/sincronizar-detalle-ventas", data={"rut": "12345678-9", "desde": "202403", "hasta": "202403"}
    )
    assert respuesta.status_code == 400


def test_detalle_ventas_acepta_fecha_exacta_aaaammdd(cliente_con_datos):
    """Antes sólo aceptaba AAAAMM; una fecha exacta como "20240301" tiraba un
    400 porque se intentaba interpretar como periodo mensual."""
    respuesta = cliente_con_datos.post(
        "/sincronizar-detalle-ventas",
        data={"rut": RUT, "desde": "20240301", "hasta": "20240315"},
        follow_redirects=False,
    )
    assert respuesta.status_code == 303


def test_detalle_ventas_fecha_inicial_posterior_a_la_final(cliente_con_datos):
    respuesta = cliente_con_datos.post(
        "/sincronizar-detalle-ventas",
        data={"rut": RUT, "desde": "20240315", "hasta": "20240301"},
    )
    assert respuesta.status_code == 400


def test_credencial_con_rut_de_acceso_distinto(cliente):
    """El representante legal (11433270-4) puede guardar la credencial de una
    empresa (76655600-0) que representa, con su propio RUT de acceso."""
    from sqlalchemy import select

    cliente.post(
        "/credenciales",
        data={"rut": "76655600-0", "rut_acceso": "11433270-4", "clave": "clave-del-representante"},
    )
    pagina = cliente.get("/credenciales")
    assert "76655600-0" in pagina.text
    assert "11433270-4" in pagina.text

    from app.db import SessionLocal
    from app.models import Credencial

    with SessionLocal() as db:
        credencial = db.scalars(select(Credencial)).one()
        assert credencial.rut_titular == "76655600-0"
        assert credencial.rut == "11433270-4"
        assert credencial.clave_cifrada != "clave-del-representante"


def test_dos_empresas_con_el_mismo_representante(cliente):
    """Una misma persona puede guardar credenciales para más de una empresa."""
    from sqlalchemy import select

    cliente.post("/credenciales", data={"rut": "76655600-0", "rut_acceso": "11433270-4", "clave": "clave1"})
    respuesta = cliente.post(
        "/credenciales",
        data={"rut": "76760089-5", "rut_acceso": "11433270-4", "clave": "clave2"},
        follow_redirects=False,
    )
    assert respuesta.status_code == 303

    from app.db import SessionLocal
    from app.models import Credencial

    with SessionLocal() as db:
        filas = db.scalars(select(Credencial)).all()
        assert {f.rut_titular for f in filas} == {"76655600-0", "76760089-5"}
        assert {f.rut for f in filas} == {"11433270-4"}


def test_credencial_sin_rut_de_acceso_usa_el_mismo_rut(cliente):
    from sqlalchemy import select

    cliente.post("/credenciales", data={"rut": RUT, "clave": "clave-propia"})

    from app.db import SessionLocal
    from app.models import Credencial

    with SessionLocal() as db:
        credencial = db.scalars(select(Credencial)).one()
        assert credencial.rut_titular == RUT
        assert credencial.rut == RUT


def test_filtro_tipo_doc_vacio_no_revienta(cliente_con_datos):
    """El <select> de "Tipo" manda tipo_doc="" cuando está en "Todos" — antes
    esto hacía que FastAPI intentara parsearlo como entero y fallara con 422."""
    respuesta = cliente_con_datos.get("/documentos", params={"titular": RUT, "tipo_doc": ""})
    assert respuesta.status_code == 200

    respuesta = cliente_con_datos.get("/exportar.xlsx", params={"titular": RUT, "tipo_doc": ""})
    assert respuesta.status_code == 200

    respuesta = cliente_con_datos.get("/exportar.csv", params={"titular": RUT, "tipo_doc": ""})
    assert respuesta.status_code == 200


def test_filtro_tipo_doc_con_valor_sigue_funcionando(cliente_con_datos):
    respuesta = cliente_con_datos.get("/documentos", params={"titular": RUT, "tipo_doc": "33"})
    assert respuesta.status_code == 200


def test_elimina_un_registro_del_historial_de_descargas(cliente_con_datos):
    """Borrar del historial (p. ej. un intento fallido) no debe tocar los
    documentos que ya se guardaron en otros trabajos."""
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import Sincronizacion

    with SessionLocal() as db:
        trabajo = db.scalars(select(Sincronizacion)).one()
        trabajo_id = trabajo.id

    respuesta = cliente_con_datos.post(f"/sincronizaciones/{trabajo_id}/eliminar", follow_redirects=False)
    assert respuesta.status_code == 303

    with SessionLocal() as db:
        assert db.scalars(select(Sincronizacion)).all() == []

    datos = cliente_con_datos.get("/api/documentos", params={"titular": RUT}).json()
    assert datos["total"] > 0


def test_elimina_un_trabajo_atascado_en_corriendo(cliente_con_datos):
    """Un trabajo que quedó en "corriendo" (p. ej. porque el servidor se cerró
    a mitad de una descarga) también debe poder borrarse del historial."""
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import Sincronizacion

    with SessionLocal() as db:
        trabajo = db.scalars(select(Sincronizacion)).one()
        trabajo.estado = "corriendo"
        db.commit()
        trabajo_id = trabajo.id

    respuesta = cliente_con_datos.post(f"/sincronizaciones/{trabajo_id}/eliminar", follow_redirects=False)
    assert respuesta.status_code == 303

    with SessionLocal() as db:
        assert db.scalars(select(Sincronizacion)).all() == []


def test_paginas_visibles_lista_completa_cuando_son_pocas():
    from app.web.main import _paginas_visibles

    assert _paginas_visibles(1, 5) == [1, 2, 3, 4, 5]


def test_paginas_visibles_usa_relleno_cuando_son_muchas():
    from app.web.main import _paginas_visibles

    assert _paginas_visibles(10, 20) == [1, None, 9, 10, 11, None, 20]


def test_documentos_pagina_dos_es_navegable_por_numero(cliente_con_datos):
    """Antes sólo había "Anterior"/"Siguiente"; ahora el número de página 2
    también debe aparecer como un enlace navegable."""
    cliente_con_datos.post(
        "/sincronizar",
        data={"rut": RUT, "desde": "202405", "hasta": "202412", "compras": "on", "ventas": "on"},
    )
    respuesta = cliente_con_datos.get("/documentos", params={"titular": RUT})
    assert 'href="/documentos?' in respuesta.text
    assert "pagina=2" in respuesta.text


def test_elimina_todos_los_documentos_de_un_titular(cliente_con_datos):
    """Para limpiar datos de prueba o de un RUT equivocado, sin tocar la
    credencial guardada ni el historial de Descargas."""
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import Documento, Sincronizacion

    with SessionLocal() as db:
        assert db.scalars(select(Documento).where(Documento.rut_titular == RUT)).first() is not None

    respuesta = cliente_con_datos.post(
        "/documentos/eliminar-todo", data={"titular": RUT}, follow_redirects=False
    )
    assert respuesta.status_code == 303

    with SessionLocal() as db:
        assert db.scalars(select(Documento).where(Documento.rut_titular == RUT)).all() == []
        # No toca el historial de descargas ni (por extensión) la credencial.
        assert db.scalars(select(Sincronizacion)).first() is not None


def test_elimina_solo_las_compras_de_un_titular(cliente_con_datos):
    from sqlalchemy import select

    from app.db import SessionLocal
    from app.models import Documento

    respuesta = cliente_con_datos.post(
        "/documentos/eliminar-todo",
        data={"titular": RUT, "operacion": "COMPRA"},
        follow_redirects=False,
    )
    assert respuesta.status_code == 303

    with SessionLocal() as db:
        restantes = db.scalars(select(Documento).where(Documento.rut_titular == RUT)).all()
        assert restantes
        assert all(d.operacion == "VENTA" for d in restantes)


def test_eliminar_documentos_vuelve_a_reportes_si_se_pide(cliente_con_datos):
    respuesta = cliente_con_datos.post(
        "/documentos/eliminar-todo",
        data={"titular": RUT, "volver": "/reportes"},
        follow_redirects=False,
    )
    assert respuesta.status_code == 303
    assert respuesta.headers["location"].startswith("/reportes")


def test_recuerda_el_titular_al_cambiar_de_pagina(cliente_con_datos):
    """Sin RUT en la URL, cada página debería quedarse en la última empresa
    consultada (cookie) en vez de volver siempre a la primera alfabética."""
    OTRO_RUT = "76655600-0"
    assert OTRO_RUT > RUT  # RUT sería el "primero" alfabético por defecto

    cliente_con_datos.post(
        "/sincronizar",
        data={"rut": OTRO_RUT, "desde": "202403", "hasta": "202404", "compras": "on", "ventas": "on"},
    )

    respuesta = cliente_con_datos.get("/documentos", params={"titular": OTRO_RUT})
    assert respuesta.status_code == 200
    assert f"titular_activo={OTRO_RUT}" in respuesta.headers.get("set-cookie", "")

    # Sin titular en la URL (como al hacer clic en el menú de navegación),
    # debe quedarse en OTRO_RUT gracias a la cookie, no volver a RUT.
    respuesta = cliente_con_datos.get("/reportes")
    assert OTRO_RUT in respuesta.text
    respuesta = cliente_con_datos.get("/")
    assert OTRO_RUT in respuesta.text


def test_recuerda_el_filtro_de_operacion_al_volver_sin_parametros(cliente_con_datos):
    """Filtrar por Ventas y luego entrar a /documentos sin parámetros (como
    al hacer clic en el menú) debería seguir mostrando sólo Ventas."""
    respuesta = cliente_con_datos.get("/documentos", params={"titular": RUT, "operacion": "VENTA"})
    assert respuesta.status_code == 200
    assert "filtros_documentos=" in respuesta.headers.get("set-cookie", "")

    respuesta = cliente_con_datos.get("/documentos")
    assert 'value="VENTA" selected' in respuesta.text
    assert "etiqueta-compra" not in respuesta.text.split("<tbody>")[1]


def test_limpiar_filtros_borra_lo_guardado(cliente_con_datos):
    """ "Limpiar" (?titular=X sin más parámetros) debe ganarle a lo guardado,
    y ese "sin filtro" debe quedar como lo nuevo por defecto."""
    cliente_con_datos.get("/documentos", params={"titular": RUT, "operacion": "VENTA"})

    respuesta = cliente_con_datos.get("/documentos", params={"titular": RUT})
    assert respuesta.status_code == 200

    respuesta = cliente_con_datos.get("/documentos")
    cuerpo = respuesta.text.split("<tbody>")[1]
    assert "etiqueta-compra" in cuerpo
    assert "etiqueta-venta" in cuerpo
